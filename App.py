import tkinter as tk
from tkinter import messagebox, ttk, filedialog
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import numpy as np
import openpyxl
from collections import defaultdict
import subprocess
import os
import platform

# Import Logic
from logic import growth_rate_fit, calculate_doubling_time, find_best_growth_phase, od_to_cfu_estimate, calculate_cfu_from_plate

# --- Global Data Structure ---
DATA_SERIES = defaultdict(list) 

class GrowthApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Microbial Growth Analyzer - Pro Version")
        
        # 1. Enlarge Fonts globally
        self.configure_styles()
        
        self.setup_ui()
        
    def configure_styles(self):
        style = ttk.Style()
        style.theme_use('clam') # Usually looks better for styling
        style.configure('.', font=('Segoe UI', 11))
        style.configure('Treeview', rowheight=28, font=('Segoe UI', 10))
        style.configure('Treeview.Heading', font=('Segoe UI', 11, 'bold'))
        
    def setup_ui(self):
        main_frame = ttk.PanedWindow(self.root, orient=tk.HORIZONTAL)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        left_panel = ttk.Frame(main_frame, width=450)
        right_panel = ttk.Frame(main_frame)
        main_frame.add(left_panel)
        main_frame.add(right_panel)

        # --- LEFT PANEL: Controls ---
        
        # 1. Experiment Settings
        settings_frame = ttk.LabelFrame(left_panel, text="1. Experiment Settings")
        settings_frame.pack(fill="x", pady=5)
        
        # Series Name
        ttk.Label(settings_frame, text="Series Name:").grid(row=0, column=0, sticky="w", padx=5)
        self.entry_name = ttk.Entry(settings_frame)
        self.entry_name.insert(0, "WT Glucose")
        self.entry_name.grid(row=0, column=1, sticky="ew", padx=5, pady=2)
        
        # Time Units Selection
        ttk.Label(settings_frame, text="Time Unit:").grid(row=1, column=0, sticky="w", padx=5)
        self.time_unit_var = tk.StringVar(value="Hours")
        self.combo_units = ttk.Combobox(settings_frame, textvariable=self.time_unit_var, 
                                      values=["Hours", "Minutes", "Days"], state="readonly", width=10)
        self.combo_units.grid(row=1, column=1, sticky="w", padx=5, pady=2)

        # Measurement Type (Radio)
        self.measure_type = tk.StringVar(value="OD")
        self.measure_type.trace("w", self.update_input_fields)

        type_frame = ttk.Frame(settings_frame)
        type_frame.grid(row=2, column=0, columnspan=2, sticky="ew", padx=5, pady=5)
        ttk.Radiobutton(type_frame, text="OD (Optical Density)", variable=self.measure_type, value="OD").pack(side="left", padx=5)
        ttk.Radiobutton(type_frame, text="Plate Count (CFU)", variable=self.measure_type, value="CFU").pack(side="left", padx=5)

        # Dynamic Parameters Frame
        self.params_frame = ttk.Frame(settings_frame)
        self.params_frame.grid(row=3, column=0, columnspan=2, sticky="ew", padx=5, pady=5)

        # 2. Data Entry
        entry_frame = ttk.LabelFrame(left_panel, text="2. Add / Edit Data")
        entry_frame.pack(fill="x", pady=5)
        
        # Time Input
        row1 = ttk.Frame(entry_frame)
        row1.pack(fill="x", padx=5, pady=2)
        ttk.Label(row1, text="Time:").pack(side="left")
        self.entry_t = ttk.Entry(row1, width=8)
        self.entry_t.pack(side="left", padx=5)

        # Dynamic Value Inputs
        self.input_grid = ttk.Frame(entry_frame)
        self.input_grid.pack(fill="x", padx=5, pady=2)

        btn_grid = ttk.Frame(entry_frame)
        btn_grid.pack(fill="x", padx=5, pady=5)
        ttk.Button(btn_grid, text="Add Point", command=self.add_point).pack(side="left", fill="x", expand=True, padx=2)
        ttk.Button(btn_grid, text="Remove Selected", command=self.remove_point).pack(side="left", fill="x", expand=True, padx=2)
        
        ttk.Button(entry_frame, text="Load from Excel", command=self.load_file).pack(fill="x", padx=5, pady=2)

        # 3. Actions
        action_frame = ttk.LabelFrame(left_panel, text="3. Analyze & Export")
        action_frame.pack(fill="x", pady=5)
        ttk.Button(action_frame, text="Calculate & Plot", command=self.run_analysis).pack(fill="x", padx=5, pady=5)
        ttk.Button(action_frame, text="Clear All", command=self.clear_all).pack(fill="x", padx=5, pady=2)
        ttk.Button(action_frame, text="Export Graph & Report", command=self.export_report).pack(fill="x", padx=5, pady=2)
        
        # 4. Data List
        tree_frame = ttk.Frame(left_panel)
        tree_frame.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(tree_frame, columns=("series", "t", "raw", "calc"), show="headings")
        self.tree.heading("series", text="Series")
        self.tree.heading("t", text="Time")
        self.tree.heading("raw", text="Input Data")
        self.tree.heading("calc", text="Calc. (CFU/ml)")
        
        self.tree.column("series", width=80)
        self.tree.column("t", width=50)
        self.tree.column("raw", width=120)
        self.tree.column("calc", width=120)
        
        self.tree.pack(side="left", fill="both", expand=True)
        sb = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        sb.pack(side="right", fill="y")
        self.tree.configure(yscroll=sb.set)
        
        # --- RIGHT PANEL: Visualization ---
        self.fig = plt.figure(figsize=(7, 6), dpi=100)
        self.ax = self.fig.add_subplot(111)
        self.fig.subplots_adjust(bottom=0.3) # Make room for table in export
        self.canvas = FigureCanvasTkAgg(self.fig, master=right_panel)
        self.canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        
        # Results Table (Displayed below graph in GUI)
        self.res_tree = ttk.Treeview(right_panel, columns=("name", "k", "td", "r2"), show="headings", height=6)
        self.res_tree.heading("name", text="Series")
        self.res_tree.heading("k", text="k (gen/time)")
        self.res_tree.heading("td", text="Doubling Time")
        self.res_tree.heading("r2", text="R²")
        self.res_tree.pack(fill="x", pady=5)

        # Initialize Inputs
        self.entry_blank = None
        self.entry_factor = None
        self.entry_vol = None
        self.entry_val = None
        self.entry_dil = None
        self.update_input_fields()

    def update_input_fields(self, *args):
        for widget in self.params_frame.winfo_children(): widget.destroy()
        for widget in self.input_grid.winfo_children(): widget.destroy()

        mode = self.measure_type.get()

        if mode == "OD":
            ttk.Label(self.params_frame, text="Blank OD:").grid(row=0, column=0)
            self.entry_blank = ttk.Entry(self.params_frame, width=8)
            self.entry_blank.insert(0, "0.0")
            self.entry_blank.grid(row=0, column=1, padx=5)
            
            ttk.Label(self.params_frame, text="Est. Factor (x10^8):").grid(row=0, column=2)
            self.entry_factor = ttk.Entry(self.params_frame, width=8)
            self.entry_factor.insert(0, "8.0")
            self.entry_factor.grid(row=0, column=3, padx=5)

            ttk.Label(self.input_grid, text="OD Value:").pack(side="left")
            self.entry_val = ttk.Entry(self.input_grid, width=10)
            self.entry_val.pack(side="left", padx=5)
            self.entry_dil = None 
        else:
            ttk.Label(self.params_frame, text="Plated Vol (ml):").grid(row=0, column=0)
            self.entry_vol = ttk.Entry(self.params_frame, width=8)
            self.entry_vol.insert(0, "0.01")
            self.entry_vol.grid(row=0, column=1, padx=5)

            ttk.Label(self.input_grid, text="# Colonies:").pack(side="left")
            self.entry_val = ttk.Entry(self.input_grid, width=8)
            self.entry_val.pack(side="left", padx=2)
            
            ttk.Label(self.input_grid, text="Dilution (e.g. 1000):").pack(side="left")
            self.entry_dil = ttk.Entry(self.input_grid, width=8)
            self.entry_dil.insert(0, "1000")
            self.entry_dil.pack(side="left", padx=2)

    def check_data_consistency(self, new_type):
        """Ensures all series in the list are of the same type (OD or CFU)."""
        if not DATA_SERIES: return True
        
        # Check the first point of the first series
        first_series = next(iter(DATA_SERIES.values()))
        if not first_series: return True
        
        existing_type = first_series[0]['type']
        if existing_type != new_type:
            messagebox.showerror("Type Error", 
                                 f"Cannot mix data types!\nCurrent data is {existing_type}, you tried to add {new_type}.\nPlease clear all data first.")
            return False
        return True

    def add_point(self):
        try:
            name = self.entry_name.get().strip()
            if not name: raise ValueError("Series Name required")
            
            t = float(self.entry_t.get())
            mode = self.measure_type.get()
            
            if not self.check_data_consistency(mode): return

            if mode == "OD":
                val = float(self.entry_val.get())
                if val <= 0: raise ValueError("OD must be > 0")
                DATA_SERIES[name].append({"t": t, "type": "OD", "od": val})
            else:
                count = float(self.entry_val.get())
                dil = float(self.entry_dil.get())
                if count < 0: raise ValueError("Count must be >= 0")
                DATA_SERIES[name].append({"t": t, "type": "CFU", "count": count, "dil": dil})

            self.update_data_table()
            self.entry_val.delete(0, tk.END)
        except ValueError as e:
            messagebox.showerror("Input Error", str(e))

    def remove_point(self):
        selected_item = self.tree.selection()
        if not selected_item:
            messagebox.showwarning("Select Point", "Please select a row to delete.")
            return
        
        # Get values from tree
        item = self.tree.item(selected_item)
        vals = item['values']
        series_name = vals[0]
        time_val = float(vals[1])
        
        # Find and remove from DATA_SERIES
        if series_name in DATA_SERIES:
            # We filter out the matching point (assuming distinct times per series for simplicity, or first match)
            # A robust way is to find index.
            points = DATA_SERIES[series_name]
            for i, p in enumerate(points):
                if abs(p['t'] - time_val) < 1e-6: # Float comparison
                    points.pop(i)
                    break
            if not points: del DATA_SERIES[series_name]
            
            self.update_data_table()

    def update_data_table(self):
        for i in self.tree.get_children(): self.tree.delete(i)
        
        try:
            if self.measure_type.get() == "OD":
                blank = float(self.entry_blank.get())
                factor = float(self.entry_factor.get()) * 1e8
                vol = 0
            else:
                vol = float(self.entry_vol.get())
                blank, factor = 0, 0
        except:
            blank, factor, vol = 0, 0, 0.01

        for name, points in DATA_SERIES.items():
            for p in points:
                calc_val = 0
                if p["type"] == "OD":
                    raw = f"OD: {p['od']}"
                    calc_val = od_to_cfu_estimate(p['od'], blank, factor)
                else:
                    raw = f"Cols: {int(p['count'])} (x{int(p['dil'])})"
                    try:
                        calc_val = calculate_cfu_from_plate(p['count'], p['dil'], vol)
                    except: calc_val = 0
                
                self.tree.insert('', 'end', values=(name, p['t'], raw, f"{calc_val:.2e}"))

    def clear_all(self):
        DATA_SERIES.clear()
        self.update_data_table()
        self.ax.clear()
        self.canvas.draw()
        for i in self.res_tree.get_children(): self.res_tree.delete(i)

    def load_file(self):
        filepath = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if not filepath: return
        name = self.entry_name.get()
        mode = self.measure_type.get()
        
        if not self.check_data_consistency(mode): return

        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or row[0] is None: continue
                t = float(row[0])
                if mode == "OD" and len(row) >= 2:
                    val = float(row[1])
                    DATA_SERIES[name].append({"t": t, "type": "OD", "od": val})
                elif mode == "CFU" and len(row) >= 3:
                    count = float(row[1])
                    dil = float(row[2])
                    DATA_SERIES[name].append({"t": t, "type": "CFU", "count": count, "dil": dil})
            self.update_data_table()
            messagebox.showinfo("Success", f"Loaded data into '{name}'")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def run_analysis(self):
        self.ax.clear()
        for i in self.res_tree.get_children(): self.res_tree.delete(i)
        
        unit = self.time_unit_var.get()
        mode = self.measure_type.get() # For Y-axis label

        # Update Tree headers with units
        self.res_tree.heading("k", text=f"k (gen/{unit})")
        self.res_tree.heading("td", text=f"Doubling Time ({unit})")

        try:
            if mode == "OD":
                od_blank = float(self.entry_blank.get())
                od_factor = float(self.entry_factor.get()) * 1e8
                cfu_vol = 0
            else:
                cfu_vol = float(self.entry_vol.get())
                od_blank, od_factor = 0, 0
        except:
            messagebox.showerror("Error", "Check Settings.")
            return

        colors = plt.rcParams['axes.prop_cycle'].by_key()['color']
        analysis_results = []

        for i, (name, points) in enumerate(DATA_SERIES.items()):
            points.sort(key=lambda x: x["t"])
            times, log_cfus = [], []
            
            for p in points:
                t = p["t"]
                cfu = 0
                if p["type"] == "OD":
                    cfu = od_to_cfu_estimate(p["od"], od_blank, od_factor)
                elif p["type"] == "CFU":
                    cfu = calculate_cfu_from_plate(p["count"], p["dil"], cfu_vol)
                
                if cfu > 0:
                    times.append(t)
                    log_cfus.append(np.log2(cfu))

            if len(times) < 2: continue

            real_cfus = [2**y for y in log_cfus]
            k, r2, (start, end) = find_best_growth_phase(times, real_cfus)
            td = calculate_doubling_time(k)
            
            analysis_results.append([name, f"{k:.3f}", f"{td:.2f}", f"{r2:.3f}"])

            color = colors[i % len(colors)]
            self.ax.scatter(times, log_cfus, color=color, alpha=0.5, label=f"{name} (Data)")
            
            fit_t = times[start:end]
            fit_log = log_cfus[start:end]
            if len(fit_t) > 1:
                slope = k
                intercept = np.mean(fit_log) - slope * np.mean(fit_t)
                x_line = np.linspace(min(fit_t), max(fit_t), 10)
                y_line = slope * x_line + intercept
                self.ax.plot(x_line, y_line, color=color, linewidth=2, label=f"{name} (Fit)")

            self.res_tree.insert('', 'end', values=(name, f"{k:.3f}", f"{td:.2f}", f"{r2:.3f}"))

        self.ax.set_xlabel(f"Time ({unit})")
        
        if mode == "OD":
             self.ax.set_ylabel("Log2(Est. CFU/ml)")
        else:
             self.ax.set_ylabel("Log2(CFU/ml)")

        self.ax.legend()
        self.ax.grid(True, alpha=0.3)
        self.canvas.draw()
        
        # Store results for export
        self.last_results = analysis_results

    def export_report(self):
        try:
            filepath = filedialog.asksaveasfilename(defaultextension=".png", filetypes=[("PNG Image", "*.png")])
            if not filepath: return

            # Embed results table into the plot for export
            if hasattr(self, 'last_results') and self.last_results:
                columns = ["Series", f"k (/{self.time_unit_var.get()})", f"Td ({self.time_unit_var.get()})", "R2"]
                table_data = [columns] + self.last_results
                
                # Create table at the bottom of the figure
                the_table = plt.table(cellText=table_data,
                                      loc='bottom',
                                      cellLoc='center',
                                      bbox=[0.0, -0.4, 1.0, 0.25]) # Adjust bbox to fit below graph
                the_table.auto_set_font_size(False)
                the_table.set_fontsize(9)
                
                # Adjust layout to make room for table
                plt.subplots_adjust(left=0.1, bottom=0.3)
            
            self.fig.savefig(filepath, bbox_inches='tight')
            
            # Reset layout for GUI view
            plt.subplots_adjust(bottom=0.1)
            if hasattr(self, 'last_results'): 
                # Clear table from GUI view (it's only for export)
                self.ax.tables.clear()
                self.canvas.draw()

            if messagebox.askyesno("Export Success", "Open file now?"):
                sys_plat = platform.system()
                if sys_plat == 'Windows': os.startfile(filepath)
                elif sys_plat == 'Darwin': subprocess.run(['open', filepath], check=True)
                else: subprocess.run(['xdg-open', filepath], check=True)
        except Exception as e:
            messagebox.showerror("Error", f"Could not save: {e}")

if __name__ == "__main__":
    root = tk.Tk()
    app = GrowthApp(root)
    root.mainloop()